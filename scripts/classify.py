"""
Stage 1: PDF → Markdown (marker-pdf)
Stage 2: Haiku classify questions
Stage 3: Screen against ma_tran_chuan
Stage 4: Generate ma_tran_de.xlsx
"""

from __future__ import annotations

import json
import time
from datetime import datetime, timezone
from pathlib import Path

import anthropic
import openpyxl
import yaml
from openpyxl.styles import Font, PatternFill
from rich.console import Console
from rich.table import Table

import os 
import torch
device = "cuda" if torch.cuda.is_available() else "cpu"
os.environ["PYTORCH_CUDA_ALLOC_CONF"] = "expandable_segments:True"

console = Console()

ROOT = Path(__file__).resolve().parent.parent
CONFIG_PATH = ROOT / "config" / "ma_tran_chuan.yaml"
RUBRIC_PATH = ROOT / "config" / "rubric.md"
FEW_SHOT_PATH = ROOT / "config" / "few_shot_classify.md"
RAW_DIR = ROOT / "question_bank" / "raw"
PROCESSED_DIR = ROOT / "question_bank" / "processed"


def load_config() -> dict:
    with open(CONFIG_PATH, encoding="utf-8") as f:
        return yaml.safe_load(f)


def convert_pdf_to_markdown(pdf_path: Path, exam_dir: Path) -> Path:
    md_path = exam_dir / "de.md"
    if md_path.exists():
        console.print(f"  [dim]de.md already exists, skipping marker[/dim]")
        return md_path
    try:
        from marker.converters.pdf import PdfConverter
        from marker.models import create_model_dict
        from marker.output import text_from_rendered
        # from marker.config.parser import ConfigParser

        console.print(f"  Converting PDF with marker...")
        config = {"batch_multiplier": 1
            # "layout_batch_size": 1,
            # "recognition_batch_size": 4,
            # "detection_batch_size": 2,   # text detection step, also early in the pipeline
            }
        models = create_model_dict(device=device)
        converter = PdfConverter(artifact_dict=models, config=config)
        rendered = converter(str(pdf_path))
        text, _, _ = text_from_rendered(rendered)
        md_path.write_text(text, encoding="utf-8")
        console.print(f"  [green]Saved de.md[/green]")
    except Exception as e:
        console.print(f"  [red]marker conversion failed: {e}[/red]")
        raise
    return md_path


def call_haiku_classify(exam_id: str, exam_markdown: str) -> dict:
    rubric_content = RUBRIC_PATH.read_text(encoding="utf-8")
    few_shot_content = FEW_SHOT_PATH.read_text(encoding="utf-8")

    client = anthropic.Anthropic()

    system = [
        {
            "type": "text",
            "text": (
                "You are classifying questions from a Vietnamese high school biology exam "
                "(format: Sinh học 2025, 3-part structure). Return ONLY a JSON object matching "
                "the schema. No markdown, no prose."
            ),
        },
        {
            "type": "text",
            "text": rubric_content + "\n\n" + few_shot_content,
            "cache_control": {"type": "ephemeral"},
        },
    ]

    user_message = (
        f"Classify all questions in this exam:\n\n{exam_markdown}\n\n"
        "Return JSON with key 'classifications' containing an array of classification objects. "
        "Each object must have: phan (P1/P2/P3), stt (int), chuong (code), muc_do (B/H/VD). "
        "For P2 items also include: y (a/b/c/d)."
    )

    for attempt in range(2):
        try:
            response = client.messages.create(
                model="claude-haiku-4-5-20251001",
                max_tokens=4000,
                system=system,
                messages=[{"role": "user", "content": user_message}],
            )

            usage = response.usage
            cache_read = getattr(usage, "cache_read_input_tokens", 0) or 0
            console.print(
                f"  [dim][Haiku] {exam_id}: {usage.input_tokens} in / "
                f"{usage.output_tokens} out (cached: {cache_read})[/dim]"
            )
            return json.loads(response.content[0].text)

        except json.JSONDecodeError as e:
            raw_path = PROCESSED_DIR / exam_id / "classify_raw_error.txt"
            raw_path.write_text(response.content[0].text, encoding="utf-8")
            console.print(f"  [yellow]JSON parse failed, saved raw to {raw_path}: {e}[/yellow]")
            raise

        except anthropic.APIError as e:
            if attempt == 0:
                console.print(f"  [yellow]API error ({e}), retrying in 10s...[/yellow]")
                time.sleep(10)
            else:
                raise

    raise RuntimeError("Classify API call failed after retries")


def compute_screen_score(classifications: list[dict], config: dict) -> tuple[float, int, bool]:
    phan_cfg = config["phan"]

    p1_items = [c for c in classifications if c["phan"] == "P1"]
    p2_items = [c for c in classifications if c["phan"] == "P2"]
    p3_items = [c for c in classifications if c["phan"] == "P3"]

    actual_p1 = len(p1_items)
    actual_p2 = len(p2_items)
    actual_p3 = len(p3_items)

    expected_p1 = phan_cfg["P1"]["so_cau"]
    expected_p2 = phan_cfg["P2"]["so_y"]
    expected_p3 = phan_cfg["P3"]["so_cau"]

    def section_dev(actual: int, expected: int) -> float:
        return abs(actual - expected) / expected if expected else 0

    section_score = 1 - sum([
        section_dev(actual_p1, expected_p1),
        section_dev(actual_p2, expected_p2),
        section_dev(actual_p3, expected_p3),
    ]) / 3

    all_muc_do_expected = {}
    for phan_key, pdata in phan_cfg.items():
        for muc, cnt in pdata.get("muc_do", {}).items():
            all_muc_do_expected[muc] = all_muc_do_expected.get(muc, 0) + cnt

    all_muc_do_actual: dict[str, int] = {}
    for c in classifications:
        md = c["muc_do"]
        all_muc_do_actual[md] = all_muc_do_actual.get(md, 0) + 1

    muc_devs = []
    for muc, expected in all_muc_do_expected.items():
        actual = all_muc_do_actual.get(muc, 0)
        muc_devs.append(abs(actual - expected) / expected if expected else 0)

    muc_do_score = 1 - (sum(muc_devs) / len(muc_devs)) if muc_devs else 1.0

    match_score = 0.6 * section_score + 0.4 * muc_do_score
    match_score = max(0.0, min(1.0, match_score))

    vd_count = sum(1 for c in classifications if c["muc_do"] == "VD")
    screen_pass = match_score >= config["screen"]["list_a_match_min"]

    return round(match_score, 4), vd_count, screen_pass


def build_matrix_excel(exam_id: str, classifications: list[dict], exam_dir: Path, config: dict):
    chapters = list(config["chuong"].keys())
    sections = ["P1", "P2", "P3"]
    muc_dos = ["B", "H", "VD"]

    counts: dict[str, dict[str, dict[str, int]]] = {}
    for ch in chapters:
        counts[ch] = {p: {m: 0 for m in muc_dos} for p in sections}

    for c in classifications:
        ch = c.get("chuong", "")
        phan = c.get("phan", "")
        md = c.get("muc_do", "")
        if ch in counts and phan in counts[ch] and md in counts[ch][phan]:
            counts[ch][phan][md] += 1

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ma_Tran"

    col_headers = []
    for p in sections:
        for m in muc_dos:
            col_headers.append(f"{p}_{m}")
    col_headers.append("Total")

    header_row = ["Chương"] + col_headers
    ws.append(header_row)

    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    ws.freeze_panes = "B2"

    for ch in chapters:
        ten = config["chuong"][ch]["ten"]
        row = [f"{ch} — {ten}"]
        total = 0
        for p in sections:
            for m in muc_dos:
                v = counts[ch][p][m]
                row.append(v)
                total += v
        row.append(total)
        ws.append(row)

    ref_row = ["Chuẩn BộGD"]
    phan_cfg = config["phan"]
    ref_total = 0
    for p in sections:
        for m in muc_dos:
            v = phan_cfg[p]["muc_do"].get(m, 0)
            ref_row.append(v)
            ref_total += v
    ref_row.append(ref_total)
    ws.append(ref_row)

    last_row = ws.max_row
    yellow = PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid")
    for cell in ws[last_row]:
        cell.font = bold
        cell.fill = yellow

    ws.column_dimensions["A"].width = 40
    for col_letter in "BCDEFGHIJKL":
        ws.column_dimensions[col_letter].width = 8

    out_path = exam_dir / "ma_tran_de.xlsx"
    wb.save(out_path)
    console.print(f"  [green]Saved ma_tran_de.xlsx[/green]")


def classify_exam(pdf_path: Path, force: bool = False) -> dict | None:
    exam_id = pdf_path.stem
    exam_dir = PROCESSED_DIR / exam_id
    exam_dir.mkdir(parents=True, exist_ok=True)

    classify_json_path = exam_dir / "classify.json"
    if classify_json_path.exists() and not force:
        console.print(f"[yellow]Skipping {exam_id} (already classified, use --force to reprocess)[/yellow]")
        with open(classify_json_path, encoding="utf-8") as f:
            return json.load(f)

    console.print(f"\n[bold cyan]Processing {exam_id}[/bold cyan]")

    try:
        md_path = convert_pdf_to_markdown(pdf_path, exam_dir)
    except Exception as e:
        console.print(f"  [red]Skipping {exam_id}: marker failed ({e})[/red]")
        return None

    exam_markdown = md_path.read_text(encoding="utf-8")

    try:
        result = call_haiku_classify(exam_id, exam_markdown)
    except Exception as e:
        console.print(f"  [red]Skipping {exam_id}: classify failed ({e})[/red]")
        return None

    classifications = result.get("classifications", [])

    config = load_config()
    match_score, vd_count, screen_pass = compute_screen_score(classifications, config)

    p1_actual = len([c for c in classifications if c["phan"] == "P1"])
    p2_actual = len([c for c in classifications if c["phan"] == "P2"])
    p3_actual = len([c for c in classifications if c["phan"] == "P3"])

    muc_do_counts: dict[str, int] = {}
    for c in classifications:
        md = c["muc_do"]
        muc_do_counts[md] = muc_do_counts.get(md, 0) + 1

    output = {
        "exam_id": exam_id,
        "source_pdf": str(pdf_path.relative_to(ROOT)),
        "ma_tran_match_score": match_score,
        "vd_count": vd_count,
        "section_counts": {"P1": p1_actual, "P2_items": p2_actual, "P3": p3_actual},
        "muc_do_counts": muc_do_counts,
        "classifications": classifications,
        "screen_pass": screen_pass,
        "processed_at": datetime.now(timezone.utc).isoformat(),
    }

    classify_json_path.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    console.print(f"  [green]Saved classify.json[/green] (match={match_score:.2f}, vd={vd_count}, pass={screen_pass})")

    build_matrix_excel(exam_id, classifications, exam_dir, config)

    return output


def run_classify_all(force: bool = False) -> list[dict]:
    pdf_files = list(RAW_DIR.glob("*.pdf"))
    if not pdf_files:
        console.print("[yellow]No PDFs found in question_bank/raw/[/yellow]")
        return []

    results = []
    from rich.progress import Progress

    with Progress(console=console) as progress:
        task = progress.add_task("Classifying exams...", total=len(pdf_files))
        for pdf_path in pdf_files:
            exam_id = pdf_path.stem
            classify_json = PROCESSED_DIR / exam_id / "classify.json"
            if classify_json.exists() and not force:
                progress.advance(task)
                continue
            result = classify_exam(pdf_path, force=force)
            if result:
                results.append(result)
            progress.advance(task)

    return results


def print_summary_table(results: list[dict]):
    table = Table(title="Classification Results", show_header=True)
    table.add_column("Exam ID", style="cyan")
    table.add_column("Match Score", justify="right")
    table.add_column("VD Count", justify="right")
    table.add_column("Screen Pass", justify="center")

    for r in results:
        pass_str = "[green]✓[/green]" if r.get("screen_pass") else "[red]✗[/red]"
        table.add_row(
            r["exam_id"],
            f"{r['ma_tran_match_score']:.3f}",
            str(r["vd_count"]),
            pass_str,
        )

    console.print(table)
