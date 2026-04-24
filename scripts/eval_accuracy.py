"""
Compare classify output against eval/ground_truth.xlsx for accuracy metrics.
"""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

import openpyxl
import yaml
from rich.console import Console
from rich.table import Table

console = Console()

ROOT = Path(__file__).resolve().parent.parent
CONFIG_PATH = ROOT / "config" / "ma_tran_chuan.yaml"
PROCESSED_DIR = ROOT / "question_bank" / "processed"
EVAL_DIR = ROOT / "eval"
GROUND_TRUTH_PATH = EVAL_DIR / "ground_truth.xlsx"
REPORT_PATH = EVAL_DIR / "accuracy_report.md"


def load_config() -> dict:
    with open(CONFIG_PATH, encoding="utf-8") as f:
        return yaml.safe_load(f)


def load_ground_truth() -> list[dict]:
    if not GROUND_TRUTH_PATH.exists():
        raise FileNotFoundError(
            f"Ground truth not found at {GROUND_TRUTH_PATH}.\n"
            f"Copy eval/ground_truth_template.xlsx to eval/ground_truth.xlsx and fill it in."
        )

    wb = openpyxl.load_workbook(GROUND_TRUTH_PATH)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        rows.append(dict(zip(headers, row)))
    return rows


def load_classify_for_exam(exam_id: str) -> dict:
    path = PROCESSED_DIR / exam_id / "classify.json"
    if not path.exists():
        return {}
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def find_classified(classify_data: dict, phan: str, stt: int, y: str | None) -> dict | None:
    for c in classify_data.get("classifications", []):
        if c["phan"] == phan and c["stt"] == stt:
            if phan == "P2":
                if c.get("y") == y:
                    return c
            else:
                return c
    return None


def run_eval_accuracy():
    config = load_config()

    try:
        ground_truth = load_ground_truth()
    except FileNotFoundError as e:
        console.print(f"[red]{e}[/red]")
        return

    if not ground_truth:
        console.print("[yellow]Ground truth file is empty.[/yellow]")
        return

    classify_cache: dict[str, dict] = {}

    total = 0
    correct_chuong = 0
    correct_muc_do = 0
    correct_both = 0

    by_chapter: dict[str, dict[str, int]] = defaultdict(lambda: {"total": 0, "correct_chuong": 0, "correct_muc_do": 0})
    muc_do_confusion: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))

    for row in ground_truth:
        exam_id = str(row.get("exam_id", "")).strip()
        phan = str(row.get("phan", "")).strip()
        stt = int(row.get("stt", 0))
        y = str(row.get("y", "")).strip() or None
        chuong_correct = str(row.get("chuong_correct", "")).strip()
        muc_do_correct = str(row.get("muc_do_correct", "")).strip()

        if not exam_id or not phan or not chuong_correct:
            continue

        if exam_id not in classify_cache:
            classify_cache[exam_id] = load_classify_for_exam(exam_id)

        classified = find_classified(classify_cache[exam_id], phan, stt, y)
        if classified is None:
            continue

        pred_chuong = classified.get("chuong", "")
        pred_muc_do = classified.get("muc_do", "")

        total += 1
        ch_ok = pred_chuong == chuong_correct
        md_ok = pred_muc_do == muc_do_correct

        if ch_ok:
            correct_chuong += 1
        if md_ok:
            correct_muc_do += 1
        if ch_ok and md_ok:
            correct_both += 1

        by_chapter[chuong_correct]["total"] += 1
        if ch_ok:
            by_chapter[chuong_correct]["correct_chuong"] += 1
        if md_ok:
            by_chapter[chuong_correct]["correct_muc_do"] += 1

        muc_do_confusion[muc_do_correct][pred_muc_do] += 1

    if total == 0:
        console.print("[yellow]No matching records found between ground truth and classify outputs.[/yellow]")
        return

    console.print(f"\n[bold]Overall Accuracy ({total} items):[/bold]")
    console.print(f"  Chương: {correct_chuong}/{total} = {correct_chuong/total:.1%}")
    console.print(f"  Mức độ: {correct_muc_do}/{total} = {correct_muc_do/total:.1%}")
    console.print(f"  Both:   {correct_both}/{total} = {correct_both/total:.1%}")

    chapter_table = Table(title="Accuracy by Chapter")
    chapter_table.add_column("Chapter")
    chapter_table.add_column("Total", justify="right")
    chapter_table.add_column("Chương %", justify="right")
    chapter_table.add_column("Mức độ %", justify="right")

    for ch, stats in sorted(by_chapter.items()):
        t = stats["total"]
        cc = f"{stats['correct_chuong']/t:.0%}" if t else "—"
        cm = f"{stats['correct_muc_do']/t:.0%}" if t else "—"
        chapter_table.add_row(ch, str(t), cc, cm)

    console.print(chapter_table)

    muc_dos = ["B", "H", "VD"]
    confusion_table = Table(title="Mức độ Confusion Matrix (rows=actual, cols=predicted)")
    confusion_table.add_column("Actual \\ Pred")
    for m in muc_dos:
        confusion_table.add_column(m, justify="right")

    for actual in muc_dos:
        row_vals = [actual]
        for pred in muc_dos:
            row_vals.append(str(muc_do_confusion[actual].get(pred, 0)))
        confusion_table.add_row(*row_vals)

    console.print(confusion_table)

    report_lines = [
        "# Accuracy Report",
        "",
        f"Total items evaluated: {total}",
        "",
        "## Overall",
        f"- Chương accuracy: {correct_chuong}/{total} = {correct_chuong/total:.1%}",
        f"- Mức độ accuracy: {correct_muc_do}/{total} = {correct_muc_do/total:.1%}",
        f"- Both correct: {correct_both}/{total} = {correct_both/total:.1%}",
        "",
        "## By Chapter",
        "| Chapter | Total | Chương % | Mức độ % |",
        "|---------|-------|----------|----------|",
    ]
    for ch, stats in sorted(by_chapter.items()):
        t = stats["total"]
        cc = f"{stats['correct_chuong']/t:.0%}" if t else "—"
        cm = f"{stats['correct_muc_do']/t:.0%}" if t else "—"
        report_lines.append(f"| {ch} | {t} | {cc} | {cm} |")

    report_lines += ["", "## Mức độ Confusion Matrix", "| Actual \\ Pred | B | H | VD |", "|---|---|---|---|"]
    for actual in muc_dos:
        row_vals = [actual] + [str(muc_do_confusion[actual].get(pred, 0)) for pred in muc_dos]
        report_lines.append("| " + " | ".join(row_vals) + " |")

    REPORT_PATH.write_text("\n".join(report_lines) + "\n", encoding="utf-8")
    console.print(f"\n[green]Report saved to {REPORT_PATH}[/green]")
